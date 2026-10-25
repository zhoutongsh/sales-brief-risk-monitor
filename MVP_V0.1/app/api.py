"""MVP HTTP API。

统一前缀 /api；所有业务数据在服务端按角色做行级过滤。
金额在数据库为元，离开 API 前统一转万元、保留 1 位小数；缺失返回 null。
"""
from __future__ import annotations

import io
import json
import sqlite3
from datetime import date
from html import escape

from fastapi import APIRouter, Body, Depends, HTTPException, Query, Request, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from .audit import write_audit
from .auth import (
    authenticate_demo_user,
    clear_session,
    current_user,
    demo_login_options,
    get_db,
    issue_session,
    request_meta,
    require_roles,
    user_profile,
)
from .batch import run_batch
from .brief import build_weekly_report, load_daily_brief
from .config import (
    DEMO_DATE,
    LEVEL_LABELS,
    RISK_TYPES,
    STATUS_LABELS,
    now_iso,
    parse_dt,
)
from .db import transaction
from .metrics import build_drill, load_metric_payload
from .quality import load_quality_report
from .risk_service import (
    RiskServiceError,
    check_overdue_escalations,
    claim_risk,
    close_risk,
    escalate_risk,
    query_risks,
    risk_detail_data,
    transfer_risk,
)
from .rules import get_rule, load_rules

router = APIRouter(prefix="/api")


# ---------------------------------------------------------------------------
# 请求模型
# ---------------------------------------------------------------------------

class LoginRequest(BaseModel):
    user_id: str = Field(min_length=1, max_length=64)


class ClaimRequest(BaseModel):
    plan: str = ""
    due: str | None = None
    sync: bool = False


class TransferRequest(BaseModel):
    to_user_id: str = Field(min_length=1, max_length=64)
    note: str = Field(min_length=1, max_length=500)


class CloseRequest(BaseModel):
    result_status: str
    reason: str
    action: str
    result: str


class EscalateRequest(BaseModel):
    note: str = ""


class RulePatchRequest(BaseModel):
    enabled: bool | None = None
    params: dict | None = None


class BatchRequest(BaseModel):
    date: str | None = None


# ---------------------------------------------------------------------------
# 通用
# ---------------------------------------------------------------------------

def parse_date(value: str | None, default: date = DEMO_DATE) -> date:
    if not value:
        return default
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise HTTPException(400, "日期格式必须为 YYYY-MM-DD") from exc


def _audit_read(conn: sqlite3.Connection, user: dict, action: str, resource_type: str,
                resource_id: str | None, request: Request, detail: str = "") -> None:
    meta = request_meta(request)
    write_audit(conn, user["user_id"], action, resource_type, resource_id, detail,
                meta.get("ip"), meta.get("user_agent"))


def _meta(request: Request) -> dict:
    return request_meta(request)


def _notify_item(row: sqlite3.Row | dict) -> dict:
    item = dict(row)
    item["is_read"] = bool(item.get("is_read"))
    return item


# ---------------------------------------------------------------------------
# 认证与演示账号
# ---------------------------------------------------------------------------

@router.post("/auth/login")
def login(payload: LoginRequest, request: Request, response: Response,
          conn: sqlite3.Connection = Depends(get_db)):
    user = authenticate_demo_user(conn, payload.user_id)
    issue_session(response, user["user_id"])
    meta = request_meta(request)
    write_audit(conn, user["user_id"], "auth.login", "session", user["user_id"],
                "演示账号登录", meta.get("ip"), meta.get("user_agent"))
    return {"user": user_profile(conn, user), "notice": "演示环境：非生产认证"}


@router.post("/auth/logout")
def logout(request: Request, response: Response, conn: sqlite3.Connection = Depends(get_db)):
    try:
        user = current_user(request, conn)
    except HTTPException:
        user = None
    if user:
        meta = request_meta(request)
        write_audit(conn, user["user_id"], "auth.logout", "session", user["user_id"],
                    "退出登录", meta.get("ip"), meta.get("user_agent"))
    clear_session(response)
    return {"ok": True}


@router.get("/me")
def me(request: Request, conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    return user_profile(conn, user)


@router.get("/demo/users")
def demo_accounts(conn: sqlite3.Connection = Depends(get_db)):
    items = demo_login_options(conn)
    return {
        "items": items,
        "total": len(items),
        "demo_mode": True,
        "notice": "演示环境：只允许以下预置账号登录，不是生产认证；所有业务数据均为本地演示数据。",
        "server_time": now_iso(),
    }


# ---------------------------------------------------------------------------
# 日报 / 周报
# ---------------------------------------------------------------------------

@router.get("/brief/daily")
def daily_brief(request: Request, date_value: str | None = Query(None, alias="date"),
                conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    d = parse_date(date_value)
    payload = load_daily_brief(conn, user, d)
    if not payload:
        raise HTTPException(404, "该日期暂无已发布日报，请先运行批处理")
    _audit_read(conn, user, "brief.view", "daily_brief", d.isoformat(), request,
                f"查看日报；批次 {payload.get('batch_id') or '—'}")
    result = dict(payload)
    result["server_time"] = now_iso()
    result["data_status"] = "published"
    return result


@router.get("/brief/weekly")
def weekly_brief(request: Request, week: str = Query("this"),
                 date_value: str | None = Query(None, alias="date"),
                 conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    week_type = week if week in {"this", "last"} else "this"
    d = parse_date(date_value)
    row = conn.execute(
        """SELECT * FROM weekly_reports WHERE role=? AND user_id=? AND week_type=?
           ORDER BY id DESC LIMIT 1""",
        (user["role"], user["user_id"], week_type),
    ).fetchone()
    if not row:
        raise HTTPException(404, "该周暂无已发布周报，请先运行批处理")
    # 周报闭环状态应随风险处理实时变化：API 时点重算，使用批处理发布的批次与时间。
    stored_payload = json.loads(row["payload_json"]) if row["payload_json"] else {}
    payload = build_weekly_report(
        conn, user, d, week_type, batch_id=row["batch_id"], published_at=row["published_at"],
        generated_at=stored_payload.get("generated_at"),
    )
    _audit_read(conn, user, "brief.view", "weekly_report",
                f"{week_type}:{payload['week_start']}", request, "查看周报")
    payload["server_time"] = now_iso()
    payload["source"] = "live_recalculation"
    return payload


# ---------------------------------------------------------------------------
# 指标与下钻
# ---------------------------------------------------------------------------

@router.get("/metrics")
def metrics(request: Request, date_value: str | None = Query(None, alias="date"),
            conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    d = parse_date(date_value)
    payload = load_metric_payload(conn, user, d)
    _audit_read(conn, user, "metric.view", "metric", d.isoformat(), request,
                f"查看指标；范围 {payload['scope']['name']}")
    payload["server_time"] = now_iso()
    return payload


@router.get("/metrics/{metric_key}/drill")
def metric_drill(metric_key: str, request: Request,
                 date_value: str | None = Query(None, alias="date"),
                 conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    d = parse_date(date_value)
    payload = build_drill(conn, user, metric_key, d)
    if not payload:
        raise HTTPException(404, "指标不存在")
    _audit_read(conn, user, "metric.drill", "metric", metric_key, request,
                f"下钻指标 {metric_key}")
    payload["server_time"] = now_iso()
    return payload


# ---------------------------------------------------------------------------
# 风险
# ---------------------------------------------------------------------------

def _refresh_overdue_risks(conn: sqlite3.Connection) -> None:
    """打开风险列表/详情时惰性检查超时升级，保证不依赖当天是否重跑批处理。"""
    with transaction(conn):
        check_overdue_escalations(conn)


@router.get("/risks")
def risks_list(
    request: Request,
    level: str | None = None,
    status: str | None = None,
    type: str | None = None,
    owner: str | None = None,
    q: str | None = None,
    sort: str = "priority",
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    conn: sqlite3.Connection = Depends(get_db),
):
    user = current_user(request, conn)
    _refresh_overdue_risks(conn)
    result = query_risks(conn, user, level, status, type, owner, q, sort, page, page_size)
    _audit_read(conn, user, "risk.list", "risk", None, request,
                f"风险列表 level={level or 'all'} status={status or 'all'} q={q or ''}")
    result["server_time"] = now_iso()
    return result


@router.get("/risks/{risk_id}")
def risk_detail(risk_id: str, request: Request, conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    _refresh_overdue_risks(conn)
    data = risk_detail_data(conn, user, risk_id)
    _audit_read(conn, user, "risk.view", "risk", risk_id, request, "查看风险详情")
    data["server_time"] = now_iso()
    return data


@router.post("/risks/{risk_id}/claim")
def risk_claim(risk_id: str, payload: ClaimRequest, request: Request,
               conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    with transaction(conn):
        data = claim_risk(conn, user, risk_id, payload.model_dump(), _meta(request))
    data["server_time"] = now_iso()
    return data


@router.post("/risks/{risk_id}/transfer")
def risk_transfer(risk_id: str, payload: TransferRequest, request: Request,
                  conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    with transaction(conn):
        data = transfer_risk(conn, user, risk_id, payload.to_user_id, payload.note, _meta(request))
    data["server_time"] = now_iso()
    return data


@router.post("/risks/{risk_id}/close")
def risk_close(risk_id: str, payload: CloseRequest, request: Request,
               conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    with transaction(conn):
        data = close_risk(conn, user, risk_id, payload.model_dump(), _meta(request))
    data["server_time"] = now_iso()
    return data


@router.post("/risks/{risk_id}/escalate")
def risk_escalate(risk_id: str, request: Request,
                  payload: EscalateRequest | None = Body(None),
                  conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    note = (payload.note if payload else "") or "管理员/上级手动升级"
    with transaction(conn):
        data = escalate_risk(conn, user, risk_id, note, _meta(request))
    data["server_time"] = now_iso()
    return data


# ---------------------------------------------------------------------------
# 预警规则
# ---------------------------------------------------------------------------

PARAM_RULES = {
    "target_gap": {"yellow_gap_pp", "orange_gap_pp", "red_gap_pp"},
    "decline": {"orange_drop_pct", "red_drop_pct", "workdays"},
    "leads": {"volume_drop_pct", "conversion_drop_pct"},
    "coverage": {"min_coverage"},
    "churn": {"min_amount", "days_no_followup"},
    "payment": {"min_cancel_amount", "overdue_days"},
}


@router.get("/rules")
def rules_list(request: Request, conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    items = load_rules(conn)
    for item in items:
        item["editable"] = user["role"] == "gm"
        item["affects"] = "下一次批处理；历史风险不会被改写"
    _audit_read(conn, user, "rule.view", "risk_rule", None, request, "查看预警规则")
    return {"items": items, "total": len(items), "can_edit": user["role"] == "gm",
            "server_time": now_iso()}


@router.patch("/rules/{rule_id}")
def rule_patch(rule_id: str, payload: RulePatchRequest, request: Request,
               conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    require_roles(user, "gm")
    rule = get_rule(conn, rule_id)
    if not rule:
        raise HTTPException(404, "预警规则不存在")
    updates: list[str] = []
    values: list = []
    if payload.enabled is not None:
        updates.append("enabled=?")
        values.append(1 if payload.enabled else 0)
    if payload.params is not None:
        allowed = PARAM_RULES.get(rule_id, set())
        merged = dict(rule.get("params") or {})
        for key, value in payload.params.items():
            if key not in allowed:
                raise HTTPException(400, f"规则 {rule_id} 不支持参数 {key}")
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise HTTPException(400, f"参数 {key} 必须是数字")
            if value <= 0:
                raise HTTPException(400, f"参数 {key} 必须大于 0")
            merged[key] = value
        if rule_id == "target_gap":
            if not (merged.get("yellow_gap_pp", 0) < merged.get("orange_gap_pp", 0)
                    < merged.get("red_gap_pp", 0)):
                raise HTTPException(400, "目标缺口阈值必须满足 黄色 < 橙色 < 红色")
        if rule_id == "decline":
            if not (merged.get("orange_drop_pct", 0) < merged.get("red_drop_pct", 0)):
                raise HTTPException(400, "业绩下滑阈值必须满足 橙色 < 红色")
        if rule_id == "decline" and "workdays" in merged:
            if int(merged["workdays"]) != merged["workdays"]:
                raise HTTPException(400, "工作日数量必须为整数")
        updates.append("params_json=?")
        values.append(json.dumps(merged, ensure_ascii=False))
    if not updates:
        raise HTTPException(400, "请提供 enabled 或 params")
    updates.extend(["updated_at=?", "updated_by=?"])
    values.extend([now_iso(), user["user_id"]])
    with transaction(conn):
        conn.execute(f"UPDATE risk_rules SET {', '.join(updates)} WHERE rule_id=?", values + [rule_id])
        write_audit(
            conn, user["user_id"], "rule.update", "risk_rule", rule_id,
            json.dumps({"enabled": payload.enabled, "params": payload.params}, ensure_ascii=False),
            _meta(request).get("ip"), _meta(request).get("user_agent"),
        )
    updated = get_rule(conn, rule_id)
    updated["editable"] = True
    updated["affects"] = "下一次批处理；历史风险不会被改写"
    return updated


# ---------------------------------------------------------------------------
# 数据质量
# ---------------------------------------------------------------------------

@router.get("/quality")
def quality(request: Request, date_value: str | None = Query(None, alias="date"),
            conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    d = parse_date(date_value)
    report = load_quality_report(conn, d)
    if not report:
        raise HTTPException(404, "该日期暂无数据质量报告，请先运行批处理")
    _audit_read(conn, user, "quality.view", "quality_report", d.isoformat(), request, "查看数据质量")
    report["server_time"] = now_iso()
    return report


# ---------------------------------------------------------------------------
# 站内通知
# ---------------------------------------------------------------------------

@router.get("/notifications")
def notifications(request: Request, unread: bool | None = None,
                  page: int = Query(1, ge=1), page_size: int = Query(20, ge=1, le=100),
                  conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    where = "user_id=?"
    params: list = [user["user_id"]]
    if unread:
        where += " AND is_read=0"
    total = conn.execute(f"SELECT COUNT(*) FROM notifications WHERE {where}", params).fetchone()[0]
    unread_count = conn.execute(
        "SELECT COUNT(*) FROM notifications WHERE user_id=? AND is_read=0", (user["user_id"],)
    ).fetchone()[0]
    rows = conn.execute(
        f"SELECT * FROM notifications WHERE {where} ORDER BY is_read ASC, created_at DESC, id DESC "
        "LIMIT ? OFFSET ?",
        params + [page_size, (page - 1) * page_size],
    ).fetchall()
    return {
        "items": [_notify_item(row) for row in rows],
        "total": int(total), "page": page, "page_size": page_size,
        "unread": int(unread_count), "server_time": now_iso(),
    }


@router.post("/notifications/{notification_id}/read")
def notification_read(notification_id: int, request: Request,
                      conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    row = conn.execute("SELECT * FROM notifications WHERE id=?", (notification_id,)).fetchone()
    if not row or row["user_id"] != user["user_id"]:
        raise HTTPException(404, "消息不存在")
    with transaction(conn):
        conn.execute("UPDATE notifications SET is_read=1 WHERE id=?", (notification_id,))
        write_audit(conn, user["user_id"], "notification.read", "notification",
                    str(notification_id), "标记消息已读", _meta(request).get("ip"),
                    _meta(request).get("user_agent"))
    refreshed = conn.execute("SELECT * FROM notifications WHERE id=?", (notification_id,)).fetchone()
    return _notify_item(refreshed)


@router.post("/notifications/read-all")
def notifications_read_all(request: Request, conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    before = conn.execute(
        "SELECT COUNT(*) FROM notifications WHERE user_id=? AND is_read=0", (user["user_id"],)
    ).fetchone()[0]
    with transaction(conn):
        conn.execute("UPDATE notifications SET is_read=1 WHERE user_id=? AND is_read=0",
                     (user["user_id"],))
        write_audit(conn, user["user_id"], "notification.read_all", "notification", None,
                    f"全部标记已读，共 {before} 条", _meta(request).get("ip"),
                    _meta(request).get("user_agent"))
    return {"updated": int(before), "unread": 0}


# ---------------------------------------------------------------------------
# 导出与打印
# ---------------------------------------------------------------------------

THIN = Side(style="thin", color="D9DDE5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WATERMARK = "演示数据·禁止外传"


def _watermark_sheet(ws, title: str, subtitle: str) -> int:
    ws.merge_cells("A1:H1")
    ws["A1"] = f"{title}　|　{WATERMARK}"
    ws["A1"].font = Font(size=15, bold=True, color="9B1C1C")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:H2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color="555555")
    ws["A2"].alignment = Alignment(horizontal="center")
    ws["A3"] = f"生成时间：{now_iso()}　{ WATERMARK }"
    ws["A3"].font = Font(size=9, color="9B1C1C")
    return 4


def _xlsx_response(wb: Workbook, filename: str, request: Request, conn: sqlite3.Connection,
                   user: dict, resource_id: str, detail: str) -> Response:
    buffer = io.BytesIO()
    wb.save(buffer)
    _audit_read(conn, user, "export.xlsx", filename.split("_")[0], resource_id, request, detail)
    # 响应头只能是 latin-1；中文水印写在 Excel 单元格/页眉页脚里。
    headers = {"Content-Disposition": f'attachment; filename="{filename}"',
               "X-Demo-Watermark": "demo-data-do-not-distribute"}
    return Response(content=buffer.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=headers)


def _write_table(ws, start_row: int, headers: list[str], rows: list[list], widths: list[int] | None = None):
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        if widths:
            ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    for r, values in enumerate(rows, start_row + 1):
        for c, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


@router.get("/export/brief.xlsx")
def export_brief(request: Request, date_value: str | None = Query(None, alias="date"),
                 conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    d = parse_date(date_value)
    brief = load_daily_brief(conn, user, d)
    if not brief:
        raise HTTPException(404, "该日期暂无日报可导出")
    wb = Workbook()
    ws = wb.active
    ws.title = "日报"
    start = _watermark_sheet(
        ws, f"经营日报 · {d.isoformat()}",
        f"数据范围：{brief['scope']['name']}　角色：{user.get('title', user['role'])}　发布：{brief.get('published_at')}",
    )
    kpi_rows = [[k["label"], k["display"], k.get("sub") or "", k.get("delta") if k.get("delta") is not None else "—"]
                for k in brief.get("kpis", [])]
    _write_table(ws, start, ["指标", "数值", "口径/补充", "环比/进度"], kpi_rows,
                 [18, 18, 42, 18])
    row = start + len(kpi_rows) + 3
    ws.cell(row=row, column=1, value="三条结论").font = Font(bold=True, size=12)
    row += 1
    for item in brief.get("conclusions", []):
        ws.cell(row=row, column=1, value=f"{item['index']}. {item['title']}")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row=row, column=2, value=item["text"]).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 2
    ws.cell(row=row, column=1, value="风险 Top3").font = Font(bold=True, size=12)
    row += 1
    for risk in brief.get("top_risks", []):
        owner = (risk.get("owner") or {}).get("name") or "—"
        ws.cell(row=row, column=1, value=risk["id"])
        ws.cell(row=row, column=2, value=risk["level_label"])
        ws.cell(row=row, column=3, value=risk["title"])
        ws.cell(row=row, column=4, value=risk["amount"])
        ws.cell(row=row, column=5, value=owner)
        row += 1
    row += 2
    ws.cell(row=row, column=1, value="渠道表现（近 30 天）").font = Font(bold=True, size=12)
    row += 1
    channels = brief.get("channels", [])
    if isinstance(channels, dict):
        channels = channels.get("items", [])
    _write_table(ws, row, ["渠道", "线索", "有效线索", "转化率", "归因签约(万元)", "投放(万元)", "ROI(倍)"],
                 [[c.get("channel"), c.get("leads"), c.get("valid_leads"), c.get("conversion_rate"),
                   c.get("contract_amount_wan"), c.get("spend_wan"), c.get("roi")] for c in channels])
    note_row = row + len(channels) + 3
    ws.cell(row=note_row, column=1, value="数据说明").font = Font(bold=True, size=12)
    for i, note in enumerate(brief.get("data_notes", []), 1):
        ws.cell(row=note_row + i, column=1, value=f"{i}. {note}")
        ws.merge_cells(start_row=note_row + i, start_column=1, end_row=note_row + i, end_column=8)
    ws.oddFooter.center.text = WATERMARK
    ws.oddHeader.center.text = WATERMARK
    filename = f"brief_{d.strftime('%Y%m%d')}_{user['user_id']}.xlsx"
    return _xlsx_response(wb, filename, request, conn, user, d.isoformat(), "导出日报 Excel")


@router.get("/export/risks.xlsx")
def export_risks(request: Request, level: str | None = None, status: str | None = None,
                 type: str | None = None, q: str | None = None, sort: str = "priority",
                 conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    result = query_risks(conn, user, level, status, type, None, q, sort, 1, 100, max_page_size=1000)
    wb = Workbook()
    ws = wb.active
    ws.title = "风险清单"
    filters = f"级别={level or '全部'}　状态={status or '全部'}　类型={type or '全部'}　搜索={q or '—'}"
    start = _watermark_sheet(
        ws, "风险清单", f"数据范围：{user_profile(conn, user)['scope']['label']}　{ filters }　共 {result['total']} 条",
    )
    rows = []
    for risk in result["items"]:
        rows.append([
            risk["id"], risk["level_label"], risk["type_label"], risk["title"],
            (risk.get("org") or {}).get("name") or "", (risk.get("owner") or {}).get("name") or "",
            (risk.get("owner") or {}).get("phone") or "", risk["amount"], risk["status_label"],
            risk["sla"]["claim_text"], risk["sla"]["close_text"], risk["created_at"],
        ])
    _write_table(ws, start,
                 ["风险编号", "级别", "类型", "标题", "组织", "责任人", "手机（脱敏）",
                  "金额(万元)", "状态", "认领时限", "关闭时限", "生成时间"],
                 rows, [18, 9, 14, 38, 16, 10, 16, 13, 10, 18, 18, 22])
    ws.oddFooter.center.text = WATERMARK
    ws.oddHeader.center.text = WATERMARK
    filename = f"risks_{DEMO_DATE.strftime('%Y%m%d')}_{user['user_id']}.xlsx"
    return _xlsx_response(wb, filename, request, conn, user, "risks", "导出风险清单 Excel")


@router.get("/print/brief")
def print_brief(request: Request, date_value: str | None = Query(None, alias="date"),
                conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    d = parse_date(date_value)
    brief = load_daily_brief(conn, user, d)
    if not brief:
        raise HTTPException(404, "该日期暂无日报可打印")
    _audit_read(conn, user, "print.brief", "daily_brief", d.isoformat(), request, "打印版日报")
    kpi_html = "".join(
        f"<tr><td>{escape(k['label'])}</td><td><strong>{escape(k['display'])}</strong></td>"
        f"<td>{escape(k.get('sub') or '—')}</td></tr>"
        for k in brief.get("kpis", [])
    )
    conclusions = "".join(
        f"<li><strong>{escape(c['title'])}：</strong>{escape(c['text'])}</li>"
        for c in brief.get("conclusions", [])
    )
    risk_rows = "".join(
        f"<tr><td>{escape(r['level_label'])}</td><td>{escape(r['id'])}</td>"
        f"<td>{escape(r['title'])}</td><td>{r['amount']}</td>"
        f"<td>{escape((r.get('owner') or {}).get('name') or '—')}</td></tr>"
        for r in brief.get("top_risks", [])
    )
    actions = "".join(
        f"<li>{escape(a.get('text') or '')}</li>" for a in brief.get("actions", [])
    )
    channels = brief.get("channels", [])
    if isinstance(channels, dict):
        channels = channels.get("items", [])
    channel_rows = "".join(
        f"<tr><td>{escape(c.get('channel') or '—')}</td><td>{c.get('leads')}</td>"
        f"<td>{c.get('conversion_rate') if c.get('conversion_rate') is not None else '—'}%</td>"
        f"<td>{c.get('contract_amount_wan')}</td><td>{c.get('roi')}</td></tr>"
        for c in channels
    )
    notes = "".join(f"<li>{escape(n)}</li>" for n in brief.get("data_notes", []))
    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>经营日报 {escape(d.isoformat())} · 演示数据</title>
<style>
  @page {{ size: A4; margin: 16mm; }}
  body {{ font-family: "Microsoft YaHei", sans-serif; color:#1f2937; line-height:1.7; margin:0; padding:24px; }}
  .watermark {{ position:fixed; inset:0; display:flex; align-items:center; justify-content:center;
    pointer-events:none; z-index:0; font-size:64px; font-weight:800; color:rgba(155,28,28,.08);
    transform:rotate(-28deg); letter-spacing:8px; white-space:nowrap; }}
  .content {{ position:relative; z-index:1; max-width:900px; margin:0 auto; }}
  h1 {{ font-size:24px; margin:0 0 4px; }} h2 {{ font-size:16px; border-left:4px solid #1f3a5f;
    padding-left:8px; margin:24px 0 10px; }} .meta {{ color:#64748b; font-size:12px; }}
  table {{ width:100%; border-collapse:collapse; font-size:13px; }}
  th,td {{ border:1px solid #d9dde5; padding:6px 8px; text-align:left; }}
  th {{ background:#1f3a5f; color:#fff; }} .demo {{ color:#9b1c1c; font-weight:700; }}
  .print-btn {{ margin:16px 0; padding:8px 16px; background:#1f3a5f; color:#fff; border:0; border-radius:6px; cursor:pointer; }}
  @media print {{ .print-btn {{ display:none; }} body {{ padding:0; }} .watermark {{ position:fixed; }} }}
</style></head><body>
<div class="watermark">演示数据·禁止外传</div>
<div class="content">
<button class="print-btn" onclick="window.print()">打印 / 另存为 PDF</button>
<h1>经营日报 · {escape(d.isoformat())}</h1>
<div class="meta"><span class="demo">演示数据·禁止外传</span>　范围：{escape(brief['scope']['name'])}
角色：{escape(user.get('title') or user['role'])}　发布：{escape(brief.get('published_at') or '—')}</div>
<p><strong>{escape(brief.get('headline') or '')}</strong></p>
<h2>三条结论</h2><ol>{conclusions}</ol>
<h2>六项关键指标</h2>
<table><thead><tr><th>指标</th><th>数值</th><th>口径/补充</th></tr></thead><tbody>{kpi_html}</tbody></table>
<h2>风险 Top3</h2>
<table><thead><tr><th>级别</th><th>编号</th><th>标题</th><th>金额(万元)</th><th>责任人</th></tr></thead>
<tbody>{risk_rows or '<tr><td colspan="5">暂无未结束风险</td></tr>'}</tbody></table>
<h2>动作建议</h2><ul>{actions or '<li>暂无</li>'}</ul>
<h2>渠道表现（近 30 天）</h2>
<table><thead><tr><th>渠道</th><th>线索</th><th>转化率</th><th>归因签约(万元)</th><th>ROI(倍)</th></tr></thead>
<tbody>{channel_rows or '<tr><td colspan="5">暂无</td></tr>'}</tbody></table>
<h2>数据说明</h2><ul>{notes}</ul>
<p class="meta">本页由 MVP 演示系统生成。{escape(WATERMARK)}。金额单位万元、比率保留 1 位小数；缺失值显示「—」。</p>
</div></body></html>"""
    return Response(content=html, media_type="text/html; charset=utf-8")


# ---------------------------------------------------------------------------
# 批处理与审计
# ---------------------------------------------------------------------------

@router.post("/admin/batch/run")
def admin_batch_run(request: Request, payload: BatchRequest | None = Body(None),
                    conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    require_roles(user, "gm")
    d = parse_date(payload.date if payload else None)
    result = run_batch(d, trigger_type="manual", triggered_by=user["user_id"],
                       db_path=request.app.state.db_path)
    write_audit(conn, user["user_id"], "batch.run", "batch", result["batch_id"],
                f"手动重跑批处理 {d.isoformat()}；状态 {result['status']}",
                _meta(request).get("ip"), _meta(request).get("user_agent"))
    return result


@router.get("/audit")
def audit_logs(request: Request, limit: int = Query(50, ge=1, le=500),
               action: str | None = None, page: int = Query(1, ge=1),
               conn: sqlite3.Connection = Depends(get_db)):
    user = current_user(request, conn)
    require_roles(user, "gm")
    where = "1=1"
    params: list = []
    if action:
        where += " AND a.action LIKE ?"
        params.append(f"%{action}%")
    total = conn.execute(f"SELECT COUNT(*) FROM audit_logs a WHERE {where}", params).fetchone()[0]
    rows = conn.execute(
        f"""SELECT a.*, u.name AS user_name FROM audit_logs a
            LEFT JOIN users u ON a.user_id = u.user_id
            WHERE {where} ORDER BY a.created_at DESC, a.id DESC LIMIT ? OFFSET ?""",
        params + [limit, (page - 1) * limit],
    ).fetchall()
    items = []
    for row in rows:
        item = dict(row)
        items.append(item)
    return {"items": items, "total": int(total), "page": page, "page_size": limit,
            "server_time": now_iso()}
