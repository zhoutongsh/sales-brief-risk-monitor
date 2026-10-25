"""开发自测：真实 uvicorn + httpx，覆盖 32 个端到端场景。

用法（在 MVP_V0.1 目录）：
    .venv\\Scripts\\python scripts\\selftest.py

脚本会启动一个临时数据库的本地 uvicorn，跑完自动关闭并清理。
不会修改 data/mvp.db 的生产演示库。
"""
from __future__ import annotations

import csv
import io
import json
import os
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.batch import run_batch  # noqa: E402
from app.config import DEMO_DATE, SOURCE_DIR  # noqa: E402

DB_PATH = ROOT / "data" / "_selftest.db"
EXPECTED_PATH = ROOT / "data" / "expected_risks.csv"
SERVER_LOG = ROOT / "data" / "selftest_server.log"


class TestFailure(AssertionError):
    pass


def check(condition, message):
    if not condition:
        raise TestFailure(message)


def free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def wait_health(base: str, proc: subprocess.Popen, timeout: float = 40) -> None:
    deadline = time.time() + timeout
    while time.time() < deadline:
        if proc.poll() is not None:
            raise RuntimeError(f"uvicorn 提前退出，日志见 {SERVER_LOG}")
        try:
            with urllib.request.urlopen(base + "/api/health", timeout=1) as response:
                if response.status == 200:
                    return
        except Exception:
            pass
        time.sleep(0.3)
    raise TimeoutError("uvicorn 启动超时")


def start_server() -> tuple[subprocess.Popen, str, int]:
    for suffix in ("", "-wal", "-shm"):
        path = Path(str(DB_PATH) + suffix)
        if path.exists():
            path.unlink()
    port = free_port()
    env = os.environ.copy()
    env["MVP_DB_PATH"] = str(DB_PATH)
    log = open(SERVER_LOG, "w", encoding="utf-8")
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port)],
        cwd=str(ROOT), env=env, stdout=log, stderr=subprocess.STDOUT,
    )
    base = f"http://127.0.0.1:{port}"
    try:
        wait_health(base, proc)
    except Exception:
        stop_server(proc)
        log.close()
        raise
    proc._selftest_log = log  # type: ignore[attr-defined]
    return proc, base, port


def stop_server(proc: subprocess.Popen) -> None:
    subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        proc.wait(timeout=8)
    except Exception:
        proc.kill()
    log = getattr(proc, "_selftest_log", None)
    if log:
        log.close()


def client_for(base: str) -> httpx.Client:
    return httpx.Client(base_url=base, timeout=30, follow_redirects=False)


def login_as(base: str, user_id: str) -> httpx.Client:
    client = client_for(base)
    response = client.post("/api/auth/login", json={"user_id": user_id})
    check(response.status_code == 200, f"{user_id} 登录失败：{response.status_code} {response.text[:200]}")
    return client


def read_expected() -> dict[str, dict]:
    with EXPECTED_PATH.open(encoding="utf-8-sig", newline="") as fh:
        return {row["risk_id"]: row for row in csv.DictReader(fh)}


def workbook_text(data: bytes) -> tuple:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    cells = []
    for row in ws.iter_rows(values_only=True):
        cells.extend([str(v) for v in row if v is not None])
    return wb, ws, cells


# ---------------------------------------------------------------------------
# 场景
# ---------------------------------------------------------------------------

def t01_health_and_no_mobile(base):
    client = client_for(base)
    health = client.get("/api/health")
    check(health.status_code == 200, "健康检查失败")
    check(health.json()["mobile_supported"] is False, "不应声明支持移动端")
    check(client.get("/m").status_code == 404, "移动端 /m 应不存在")
    html = client.get("/").text
    check("演示环境" in html and "非生产" in html, "首页缺少演示环境标识")
    check("assets/data.js" not in html, "首页不应引入 data.js 假数据")
    check("assets/app.js" in html, "首页未加载真实前端脚本")


def t02_unauth_guards(base):
    client = client_for(base)
    for path in ["/api/me", "/api/brief/daily", "/api/risks", "/api/quality",
                 "/api/notifications", "/api/export/brief.xlsx", "/api/print/brief", "/api/audit"]:
        response = client.get(path)
        check(response.status_code == 401, f"未登录访问 {path} 应 401，实际 {response.status_code}")
    demo = client.get("/api/demo/users")
    check(demo.status_code == 200, "登录前应能读取演示账号列表")
    check(demo.json()["total"] == 4, "演示账号应为 4 个")


def t03_demo_accounts_masked(base):
    data = client_for(base).get("/api/demo/users").json()
    users = {u["user_id"]: u for u in data["items"]}
    check(set(users) == {"u_gm", "u_sd", "u_bu", "u_sm"}, f"演示账号不符：{users.keys()}")
    for uid, user in users.items():
        check("*" in user["phone"], f"{uid} 手机号未脱敏")
        check("@" in user["email"] and "***" in user["email"], f"{uid} 邮箱未脱敏")
    check(data["demo_mode"] is True, "demo_mode 标记缺失")


def t04_login_cookie(base):
    client = client_for(base)
    response = client.post("/api/auth/login", json={"user_id": "u_gm"})
    check(response.status_code == 200, "总经理登录失败")
    cookie = " | ".join(response.headers.get_list("set-cookie")).lower()
    check("httponly" in cookie, "会话 Cookie 缺少 HttpOnly")
    check("samesite=lax" in cookie, "会话 Cookie 缺少 SameSite=Lax")
    check("max-age=28800" in cookie, "会话 Cookie 有效期不是 8 小时")
    me = client.get("/api/me")
    check(me.status_code == 200, "登录后 /api/me 失败")
    profile = me.json()
    check(profile["role"] == "gm" and len(profile["scope"]["org_ids"]) == 11, "总经理范围不是全公司")
    check(profile["phone"] == "138****1001", f"总经理手机号脱敏异常：{profile['phone']}")


def t05_daily_brief(gm):
    response = gm.get("/api/brief/daily")
    check(response.status_code == 200, "日报读取失败")
    data = response.json()
    check(data["date"] == "2026-10-24", "日报业务日期应为 2026-10-24")
    check(len(data["kpis"]) == 6, "日报应有 6 张 KPI 卡")
    check(len(data["conclusions"]) == 3, "日报应有 3 条结论")
    check(len(data.get("top_risks", [])) >= 1, "日报缺少风险 Top3")
    check(isinstance(data.get("channels"), list) and len(data["channels"]) >= 1, "渠道数据应为列表")
    displays = [k["display"] for k in data["kpis"]]
    check(any("万元" in d for d in displays), "金额展示应带万元单位")
    check("目标完成率" in " ".join(k["label"] for k in data["kpis"]), "缺目标完成率")
    print("      6 KPI:", " / ".join(f"{k['label']}={k['display']}" for k in data["kpis"]))


def t06_weekly(gm):
    response = gm.get("/api/brief/weekly?week=this")
    check(response.status_code == 200, "周报读取失败")
    data = response.json()
    check(len(data["kpis"]) == 6, "周报应有 6 个 KPI")
    check(len(data["conclusions"]) == 3, "周报应有 3 条结论")
    check(data["week_label"].startswith("本周"), "周报标签异常")
    check(data["edition"] in {"周中版", "正式版"}, f"周报版本异常：{data['edition']}")
    check("closure_rate" in data["risk_closure"], "周报缺少风险闭环率")
    last = gm.get("/api/brief/weekly?week=last")
    check(last.status_code == 200 and last.json()["week_label"].startswith("上周"), "上周周报读取失败")


def t07_metrics_drill(gm):
    response = gm.get("/api/metrics")
    check(response.status_code == 200, "指标读取失败")
    data = response.json()
    keys = [m["key"] for m in data["metrics"]]
    expected = {"new_leads", "valid_leads", "lead_valid_rate", "lead_conversion_rate",
                "new_opportunities", "pipeline_amount", "opportunity_coverage", "contract_amount",
                "contract_count", "target_completion_rate", "collection_completion_rate", "marketing_roi"}
    check(set(keys) == expected, f"12 项指标不匹配：{keys}")
    check(len(data["auxiliary"]) == 5, "辅助指标应为 5 项")
    drill = gm.get("/api/metrics/contract_amount/drill?date=2026-10-24")
    check(drill.status_code == 200, "指标下钻失败")
    detail = drill.json()
    check(detail["total"]["display"].endswith("万元"), "下钻总额单位错误")
    check(len(detail["items"]) >= 1, "下钻明细为空")
    check(all("share" in item for item in detail["items"]), "下钻缺少占比")


def t08_quality(gm):
    response = gm.get("/api/quality")
    check(response.status_code == 200, "数据质量读取失败")
    data = response.json()
    check(data["status"] == "pass", f"种子数据质量应通过，实际 {data['status']}")
    check(data["score"] == 100.0, "种子数据质量分应为 100")
    check(len(data["completeness"]["by_source"]) == 10, "应有 10 个源文件完整性记录")
    steps = [s["step"] for s in data["batch"]["steps"]]
    check(steps == ["extract", "metrics", "quality", "risks", "daily", "weekly", "publish", "notify", "escalation"],
          f"批次步骤顺序异常：{steps}")
    check(data["targets"]["orgs_with_target"] == data["targets"]["orgs_total"], "组织目标覆盖不完整")
    check(data["mapping"]["count"] == 0, "种子数据不应有映射失败")


def t09_rules(gm, sm):
    response = gm.get("/api/rules")
    check(response.status_code == 200 and response.json()["can_edit"] is True, "总经理应可编辑规则")
    check(len(response.json()["items"]) == 6, "规则应 6 条")
    sm_data = sm.get("/api/rules").json()
    check(sm_data["can_edit"] is False, "部门经理应只读规则")
    check(all(item["params"] for item in sm_data["items"]), "规则参数缺失")


def t10_risk_list_expected(gm):
    response = gm.get("/api/risks?page_size=50")
    check(response.status_code == 200, "风险列表读取失败")
    data = response.json()
    check(data["total"] == 10, f"风险应为 10 条，实际 {data['total']}")
    expected = read_expected()
    got = {item["id"]: item for item in data["items"]}
    check(set(got) == set(expected), "风险编号与 expected_risks.csv 不一致")
    for risk_id, row in expected.items():
        item = got[risk_id]
        check(item["rule_id"] == row["rule_key"], f"{risk_id} 规则不一致")
        check(item["level"] == row["level"], f"{risk_id} 级别不一致")
        check((item.get("owner") or {}).get("user_id") == row["owner_user_id"], f"{risk_id} 责任人不一致")
        expected_wan = round(int(row["amount"]) / 10000, 1)
        check(abs(float(item["amount"]) - expected_wan) < 0.01, f"{risk_id} 金额不一致")
    check(data["stats"]["by_level"] == {"red": 6, "orange": 2, "yellow": 2}, "级别统计异常")
    check(data["stats"]["red_active"] == 6, "红色未结束统计异常")


def t11_row_level_scope(base):
    sd = login_as(base, "u_sd")
    bu = login_as(base, "u_bu")
    sm = login_as(base, "u_sm")
    try:
        sd_data = sd.get("/api/risks?page_size=50").json()
        bu_data = bu.get("/api/risks?page_size=50").json()
        sm_data = sm.get("/api/risks?page_size=50").json()
        check(sd_data["total"] == 10, f"销售总监应看 10 条，实际 {sd_data['total']}")
        check(bu_data["total"] == 6, f"事业部总经理应看 6 条，实际 {bu_data['total']}")
        check(sm_data["total"] == 1 and sm_data["items"][0]["id"] == "R-20261024-008",
              f"部门经理应只看本人恒远风险，实际 {[i['id'] for i in sm_data['items']]}")
        allowed = {"E100", "E110", "E120", "E130", "E140"}
        for item in bu_data["items"]:
            check(item["org"]["org_id"] in allowed, f"事业部越权看到 {item['org']['org_id']}")
        check(len(bu.get("/api/brief/daily").json()["kpis"]) == 6, "事业部日报失败")
        check(len(sm.get("/api/brief/daily").json()["kpis"]) == 6, "部门经理日报失败")
        sm_me = sm.get("/api/me").json()
        check(sm_me["scope"]["org_ids"] == ["E110"], "部门经理范围应仅 E110")
        check(sm.get("/api/brief/daily").json()["scope"]["name"] == "华东一部", "部门经理日报范围名称异常")
    finally:
        sd.close(); bu.close(); sm.close()


def t12_risk_detail(gm):
    response = gm.get("/api/risks/R-20261024-002")
    check(response.status_code == 200, "风险详情读取失败")
    data = response.json()
    check(len(data["evidence"]) >= 3, "风险详情证据链不足")
    check(len(data["timeline"]) >= 1, "风险详情时间线为空")
    check(data["owner"]["phone"].startswith("138****"), "风险详情手机号未脱敏")
    check(data["allowed_actions"]["claim"] is True, "待认领风险应可认领")
    check(data["transfer_allowed"] is True and len(data["transfer_candidates"]) >= 4, "转派候选异常")
    check(data["object"]["name"] == "华东四部", "风险对象名称异常")


def t13_claim_close(base, sm):
    claim = sm.post("/api/risks/R-20261024-008/claim", json={
        "plan": "今天 18:00 前确认恒远科技下一步签约时间", "due": "2026-10-24 18:00", "sync": True,
    })
    check(claim.status_code == 200, f"认领失败：{claim.status_code} {claim.text[:200]}")
    data = claim.json()
    check(data["status"] == "processing", "认领后状态应为处理中")
    check(data["claimed_at"] is not None, "认领时间未写入")
    check(data["sla"]["close_due_at"] is not None, "认领后应重算关闭 SLA")
    short = sm.post("/api/risks/R-20261024-008/close", json={
        "result_status": "closed", "reason": "太短", "action": "已与客户沟通", "result": "客户确认本月底签约",
    })
    check(short.status_code == 400, f"关闭字段过短应 400，实际 {short.status_code}")
    closed = sm.post("/api/risks/R-20261024-008/close", json={
        "result_status": "closed",
        "reason": "客户预算已确认，原判断的流失风险已解除",
        "action": "已与客户电话确认预算和签约时间，并同步销售总监",
        "result": "客户确认本月底签约，风险关闭并进入正常跟进",
    })
    check(closed.status_code == 200, f"关闭失败：{closed.status_code} {closed.text[:200]}")
    closed_data = closed.json()
    check(closed_data["status"] == "closed", "关闭后状态应 closed")
    check(closed_data["close"]["reason"].startswith("客户预算"), "关闭原因未保存")
    again = sm.post("/api/risks/R-20261024-008/close", json={
        "result_status": "closed", "reason": "重复关闭测试原因", "action": "重复关闭测试动作", "result": "重复关闭测试结果",
    })
    check(again.status_code == 409, f"已关闭风险再次关闭应 409，实际 {again.status_code}")
    listing = sm.get("/api/risks?page_size=50&status=all").json()
    check(listing["stats"]["active"] == 0, "部门经理已无未结束风险")
    check(listing["stats"]["by_status"]["closed"] == 1, "已关闭统计应 1")


def t14_sm_transfer_forbidden(sm):
    response = sm.post("/api/risks/R-20261024-008/transfer",
                       json={"to_user_id": "u_ea3", "note": "尝试越权转派"})
    check(response.status_code == 403, f"部门经理转派应 403，实际 {response.status_code}")
    check("无权" in response.json()["detail"], "403 提示应说明无权")


def t15_transfer_flow(gm):
    response = gm.post("/api/risks/R-20261024-001/transfer",
                       json={"to_user_id": "u_sm", "note": "自测：转派给李娜跟进华东三部目标缺口"})
    check(response.status_code == 200, f"转派失败：{response.status_code} {response.text[:200]}")
    data = response.json()
    check(data["owner"]["user_id"] == "u_sm", "转派后责任人未变化")
    check(data["status"] == "pending", "转派后状态应回到待认领")
    watcher_ids = {w["user_id"] for w in data["watchers"]}
    check("u_ea3" in watcher_ids, "原负责人应成为关注人并保留可见权")
    check(data["sla"]["claim_due_at"] is not None, "转派后认领时限未重算")


def t16_escalate(gm):
    response = gm.post("/api/risks/R-20261024-003/escalate", json={"note": "自测：需要销售总监协调资源"})
    check(response.status_code == 200, f"升级失败：{response.status_code} {response.text[:200]}")
    data = response.json()
    check(data["escalation"]["level"] == 1, f"升级级别应为 1，实际 {data['escalation']}")
    check(data["escalation"]["escalated_to"] is not None, "升级对象缺失")


def t17_rule_admin(gm, sd):
    forbidden = sd.patch("/api/rules/coverage", json={"enabled": False})
    check(forbidden.status_code == 403, f"销售总监改规则应 403，实际 {forbidden.status_code}")
    invalid = gm.patch("/api/rules/target_gap", json={"params": {"yellow_gap_pp": 20}})
    check(invalid.status_code == 400, f"非法阈值顺序应 400，实际 {invalid.status_code}")
    disabled = gm.patch("/api/rules/coverage", json={"enabled": False})
    check(disabled.status_code == 200 and disabled.json()["enabled"] is False, "停用规则失败")
    enabled = gm.patch("/api/rules/coverage", json={"enabled": True, "params": {"min_coverage": 2.5}})
    check(enabled.status_code == 200 and enabled.json()["enabled"] is True, "恢复规则失败")
    check(float(enabled.json()["params"]["min_coverage"]) == 2.5, "规则参数恢复失败")


def t18_notifications(sm):
    response = sm.get("/api/notifications?page_size=50")
    check(response.status_code == 200, "消息读取失败")
    data = response.json()
    check(data["unread"] >= 1, "转派后部门经理应有未读消息")
    assignment = [n for n in data["items"] if n["kind"] == "assignment"]
    check(assignment, "缺少转派通知")
    check(assignment[0]["related_risk_id"] == "R-20261024-001", "转派通知未关联风险")
    read = sm.post(f"/api/notifications/{assignment[0]['id']}/read")
    check(read.status_code == 200, "标记单条已读失败")
    allread = sm.post("/api/notifications/read-all")
    check(allread.status_code == 200, "全部已读失败")
    check(sm.get("/api/notifications").json()["unread"] == 0, "全部已读后未读数不为 0")


def t19_export_brief(gm):
    response = gm.get("/api/export/brief.xlsx?date=2026-10-24")
    check(response.status_code == 200, f"日报导出失败：{response.status_code}")
    check(response.headers["content-type"].startswith("application/vnd.openxmlformats"), "日报导出类型错误")
    check("attachment" in response.headers.get("content-disposition", ""), "日报导出没有附件文件名")
    check("demo-data" in response.headers.get("x-demo-watermark", ""), "导出响应缺少水印标记")
    wb, ws, cells = workbook_text(response.content)
    check("演示数据·禁止外传" in " ".join(cells), "日报 Excel 缺少水印文字")
    check(any("三条结论" in c for c in cells), "日报 Excel 缺少结论")
    check(ws.max_row > 10, "日报 Excel 行数过少")
    wb.close()


def t20_export_risks_filtered(gm, sm):
    gm_resp = gm.get("/api/export/risks.xlsx?status=all")
    check(gm_resp.status_code == 200, "总经理风险导出失败")
    wb, ws, cells = workbook_text(gm_resp.content)
    check("演示数据·禁止外传" in " ".join(cells), "风险 Excel 缺少水印")
    phones = []
    for row in ws.iter_rows(min_row=5, max_col=7, values_only=True):
        if row[0] and str(row[0]).startswith("R-"):
            phones.append(row[6] or "")
    check(len(phones) == 10, f"总经理风险导出行数应为 10，实际 {len(phones)}")
    check(all(p == "" or ("*" in str(p)) for p in phones), f"风险导出存在未脱敏手机号：{phones}")
    wb.close()

    sm_resp = sm.get("/api/export/risks.xlsx?status=all")
    check(sm_resp.status_code == 200, "部门经理风险导出失败")
    wb2, ws2, _ = workbook_text(sm_resp.content)
    sm_rows = [row for row in ws2.iter_rows(min_row=5, max_col=7, values_only=True)
               if row[0] and str(row[0]).startswith("R-")]
    check(len(sm_rows) == 2, f"部门经理应导出 2 条风险，实际 {len(sm_rows)}")
    wb2.close()
    sm_all = sm.get("/api/risks?status=all&page_size=50").json()
    check(sm_all["total"] == 2, f"部门经理可见风险应为 2，实际 {sm_all['total']}")


def t21_print(gm):
    response = gm.get("/api/print/brief?date=2026-10-24")
    check(response.status_code == 200, "打印版日报读取失败")
    html = response.text
    check("演示数据·禁止外传" in html, "打印页缺少水印")
    check("@page" in html and "A4" in html, "打印页缺少 A4 打印样式")
    check("打印 / 另存为 PDF" in html, "打印页缺少打印按钮")
    check('href="/m"' not in html, "打印页不应包含移动端入口")


def table_counts() -> dict:
    conn = sqlite3.connect(str(DB_PATH), timeout=20)
    try:
        counts = {}
        for table in ["risks", "risk_events", "notifications", "metric_snapshots",
                      "daily_briefs", "weekly_reports", "batch_runs"]:
            counts[table] = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
        counts["risk_status"] = {
            row[0]: row[1] for row in conn.execute("SELECT id, status FROM risks ORDER BY id")
        }
        counts["risk_owner"] = {
            row[0]: row[1] for row in conn.execute("SELECT id, owner_user_id FROM risks ORDER BY id")
        }
        return counts
    finally:
        conn.close()


def t22_batch_idempotent(gm):
    before = table_counts()
    results = []
    for i in range(3):
        response = gm.post("/api/admin/batch/run", json={"date": "2026-10-24"})
        check(response.status_code == 200, f"第 {i+1} 次重跑失败：{response.status_code} {response.text[:300]}")
        data = response.json()
        check(data["status"] == "success", f"第 {i+1} 次重跑批次状态异常：{data.get('error')}")
        results.append(data)
    after = table_counts()
    for table in ["risks", "risk_events", "notifications", "metric_snapshots", "daily_briefs", "weekly_reports"]:
        check(before[table] == after[table],
              f"同日重跑后 {table} 数量变化：{before[table]} → {after[table]}")
    check(before["risk_status"] == after["risk_status"], "同日重跑改写了风险状态")
    check(before["risk_owner"] == after["risk_owner"], "同日重跑改写了责任人")
    check(after["batch_runs"] - before["batch_runs"] == 3, "batch_runs 未记录 3 次重跑")
    check("R-20261024-008" in after["risk_status"] and after["risk_status"]["R-20261024-008"] == "closed",
          "同日重跑后已关闭风险被重开")
    print("      重跑前后:", {t: after[t] for t in ["risks", "risk_events", "notifications", "metric_snapshots", "daily_briefs", "weekly_reports"]})
    print("      批次耗时:", [round(r.get("duration_seconds", 0), 3) for r in results])


def t23_batch_permission(base, sm):
    response = sm.post("/api/admin/batch/run", json={"date": "2026-10-24"})
    check(response.status_code == 403, f"部门经理重跑批次应 403，实际 {response.status_code}")
    bu = login_as(base, "u_bu")
    try:
        check(bu.post("/api/admin/batch/run", json={}).status_code == 403, "事业部重跑批次应 403")
    finally:
        bu.close()


def t24_audit(gm, sm):
    check(sm.get("/api/audit").status_code == 403, "部门经理查看审计应 403")
    response = gm.get("/api/audit?limit=500")
    check(response.status_code == 200, "审计日志读取失败")
    data = response.json()
    check(data["total"] >= 15, f"审计记录过少：{data['total']}")
    actions = {item["action"] for item in data["items"]}
    required = {"auth.login", "risk.claim", "risk.close", "risk.transfer", "risk.escalate",
                "rule.update", "print.brief", "export.xlsx", "access.denied"}
    missing = required - actions
    check(not missing, f"审计缺少动作：{missing}")
    check(any(item["user_name"] for item in data["items"]), "审计缺少操作人姓名")
    denied = [i for i in data["items"] if i["action"] == "access.denied"]
    check(any("rules" in (i.get("resource_id") or "") or "batch" in (i.get("resource_id") or "")
              for i in denied), "越权审计未记录被拒绝的路径")


def t25_openapi_and_404(base):
    client = client_for(base)
    response = client.get("/openapi.json")
    check(response.status_code == 200, "OpenAPI 文档不可用")
    paths = response.json().get("paths", {})
    for path in ["/api/auth/login", "/api/brief/daily", "/api/risks", "/api/export/brief.xlsx", "/api/admin/batch/run"]:
        check(path in paths, f"OpenAPI 缺少 {path}")
    check(client.get("/api/not-a-real-endpoint").status_code == 404, "未知 API 应 404")
    health = client.get("/api/health").json()
    check(health["demo"] is True, "健康检查未标记演示环境")


def t26_report_publication_times(gm):
    daily = gm.get("/api/brief/daily").json()
    check(daily["published_at"] == "2026-10-24T08:30:00+08:00",
          f"日报发布时间应为计划 08:30，实际 {daily['published_at']}")
    check(daily.get("generated_at") and daily["generated_at"].endswith("+08:00"),
          "日报缺少实际生成时间 generated_at")
    check(daily["generated_at"] != daily["published_at"], "日报实际生成时间不应等于计划发布时间")

    this_week = gm.get("/api/brief/weekly?week=this").json()
    last_week = gm.get("/api/brief/weekly?week=last").json()
    check(this_week["published_at"] == "2026-10-24T09:00:00+08:00",
          f"本周周报发布时间应为周六 09:00，实际 {this_week['published_at']}")
    check(last_week["published_at"] == "2026-10-19T09:00:00+08:00",
          f"上周周报发布时间应为周一 09:00，实际 {last_week['published_at']}")
    for label, data in (("本周", this_week), ("上周", last_week)):
        check(data.get("generated_at") and data["generated_at"].endswith("+08:00"),
              f"{label}周报缺少实际生成时间 generated_at")

    conn = sqlite3.connect(str(DB_PATH), timeout=20)
    try:
        row = conn.execute(
            "SELECT generated_at, published_at FROM daily_briefs WHERE user_id='u_gm' ORDER BY id DESC LIMIT 1"
        ).fetchone()
        check(row is not None, "日报数据库记录缺失")
        check(row[1] == "2026-10-24T08:30:00+08:00", f"日报库内发布时间异常：{row[1]}")
        check(row[0] != row[1], "日报库内 generated_at 应保存实际生成时间")
    finally:
        conn.close()
    print("      daily/weekly published_at:", daily["published_at"],
          this_week["published_at"], last_week["published_at"])


def t27_weekly_top3_scope(gm):
    last_week = gm.get("/api/brief/weekly?week=last").json()
    for item in last_week["top_risks"]:
        check("2026-10-12" <= item["as_of"] <= "2026-10-18",
              f"上周周报 Top3 混入区间外风险 {item['id']}（as_of={item['as_of']}）")
        check(item["as_of"] != "2026-10-24", f"上周周报 Top3 混入本周风险 {item['id']}")
    this_week = gm.get("/api/brief/weekly?week=this").json()
    check(len(this_week["top_risks"]) >= 1, "本周周报 Top3 不应为空")
    for item in this_week["top_risks"]:
        check("2026-10-19" <= item["as_of"] <= "2026-10-25",
              f"本周周报 Top3 混入区间外风险 {item['id']}（as_of={item['as_of']}）")
    print("      last/this top3:", len(last_week["top_risks"]), len(this_week["top_risks"]))


def t28_evidence_consistency(gm):
    risk_id = "R-20261024-005"
    conn = sqlite3.connect(str(DB_PATH), timeout=20)
    try:
        before = conn.execute(
            "SELECT COUNT(*) FROM risk_evidence WHERE risk_id=?", (risk_id,)
        ).fetchone()[0]
        conn.execute(
            """INSERT INTO risk_evidence
               (risk_id, evidence_key, label, current_value, threshold, baseline, source, triggered_at, sort_order)
               SELECT risk_id, evidence_key, label, current_value, threshold, baseline, source, triggered_at, sort_order
               FROM risk_evidence WHERE risk_id=?""",
            (risk_id,),
        )
        conn.commit()
        after = conn.execute(
            "SELECT COUNT(*) FROM risk_evidence WHERE risk_id=?", (risk_id,)
        ).fetchone()[0]
        distinct = conn.execute(
            "SELECT COUNT(DISTINCT evidence_key) FROM risk_evidence WHERE risk_id=?", (risk_id,)
        ).fetchone()[0]
    finally:
        conn.close()
    check(after == before * 2, "证据复制准备失败")
    detail = gm.get(f"/api/risks/{risk_id}").json()
    check(detail["evidence_count"] == after,
          f"证据计数应为 {after}，实际 {detail['evidence_count']}")
    check(len(detail["evidence"]) == detail["evidence_count"],
          f"可见证据 {len(detail['evidence'])} 条与计数 {detail['evidence_count']} 不一致")
    check(len(detail.get("latest_evidence", [])) == distinct,
          "latest_evidence 去重条数异常")
    print("      evidence rows/count/latest:", len(detail["evidence"]),
          detail["evidence_count"], len(detail["latest_evidence"]))


def t29_lazy_escalation(gm):
    risk_id = "R-20261024-002"
    conn = sqlite3.connect(str(DB_PATH), timeout=20)
    try:
        conn.execute(
            "UPDATE risks SET claim_due_at=? WHERE id=?",
            ("2000-01-01T00:00:00+08:00", risk_id),
        )
        conn.commit()
    finally:
        conn.close()

    first = gm.get(f"/api/risks/{risk_id}")
    check(first.status_code == 200, "超时风险详情读取失败")
    check(first.json()["escalation"]["level"] == 1,
          f"打开详情应自动升级 1 级，实际 {first.json()['escalation']}")

    # 重复打开列表和详情，不得继续从 1 级升到 2 级。
    gm.get("/api/risks?page_size=50")
    gm.get(f"/api/risks/{risk_id}")
    repeated = gm.get(f"/api/risks/{risk_id}").json()
    check(repeated["escalation"]["level"] == 1,
          f"重复打开不得再次升级，实际 {repeated['escalation']['level']}")

    conn = sqlite3.connect(str(DB_PATH), timeout=20)
    try:
        rows = conn.execute(
            "SELECT payload_json FROM risk_events WHERE risk_id=? AND event_type='escalated'",
            (risk_id,),
        ).fetchall()
        automatic = [row for row in rows if json.loads(row[0]).get("automatic") is True]
        level = conn.execute(
            "SELECT escalation_level FROM risks WHERE id=?", (risk_id,)
        ).fetchone()[0]
    finally:
        conn.close()
    check(level == 1, f"库内升级级别应保持 1，实际 {level}")
    check(len(automatic) == 1, f"同一时限内自动升级事件应仅 1 条，实际 {len(automatic)}")

    # 另一个风险只通过“打开列表”触发，验证列表入口同样生效且幂等。
    list_risk_id = "R-20261024-006"
    conn = sqlite3.connect(str(DB_PATH), timeout=20)
    try:
        conn.execute(
            "UPDATE risks SET claim_due_at=? WHERE id=?",
            ("2000-01-01T00:00:00+08:00", list_risk_id),
        )
        conn.commit()
    finally:
        conn.close()
    for _ in range(2):
        listing = gm.get("/api/risks?page_size=50").json()
        item = next(x for x in listing["items"] if x["id"] == list_risk_id)
        check(item["escalation"]["level"] == 1,
              f"打开列表应自动升 1 级且不重复，实际 {item['escalation']['level']}")
    print("      lazy escalation list/detail:", level, "list", 1,
          "events", len(automatic))


def t30_bu_second_level_to_gm(base):
    bu = login_as(base, "u_bu")
    try:
        first = bu.post("/api/risks/R-20261024-004/escalate", json={"note": "自测：事业部先升级一级"})
        check(first.status_code == 200, f"事业部一级升级失败：{first.status_code} {first.text[:200]}")
        data = first.json()
        check(data["escalation"]["level"] == 1, "事业部一级升级级别异常")
        check((data["escalation"].get("escalated_to") or {}).get("user_id") == "u_sd",
              f"事业部一级应升级到销售总监，实际 {data['escalation'].get('escalated_to')}")
        check(data["escalation"]["in_general_manager_pool"] is False,
              "一级升级不应进入总经理池")

        second = bu.post("/api/risks/R-20261024-004/escalate", json={"note": "自测：事业部再升级一级"})
        check(second.status_code == 200, f"事业部二级升级失败：{second.status_code} {second.text[:200]}")
        data2 = second.json()
        check(data2["escalation"]["level"] == 2, "事业部二级升级级别异常")
        check((data2["escalation"].get("escalated_to") or {}).get("user_id") == "u_gm",
              f"事业部二级应升级到总经理，实际 {data2['escalation'].get('escalated_to')}")
        check(data2["escalation"]["in_general_manager_pool"] is True,
              "二级升级进入总经理池标记应为 true")
        print("      BU escalation targets: u_sd -> u_gm")
    finally:
        bu.close()


def t31_card_conclusion_rounding(base):
    sm = login_as(base, "u_sm")
    try:
        daily = sm.get("/api/brief/daily").json()
        card = next(k for k in daily["kpis"] if k["key"] == "target_completion_rate")
        conclusion = next(c for c in daily["conclusions"] if c["metric_key"] == "target_completion_rate")
        check(card["display"] != "—", "部门经理目标完成率卡片不应为 —")
        daily_display = card["display"]
        check(daily_display in conclusion["text"],
              f"日报卡片 {daily_display} 与结论小数不一致：{conclusion['text']}")
    finally:
        sm.close()
    client = login_as(base, "u_gm")
    try:
        weekly = client.get("/api/brief/weekly?week=this").json()
        card = next(k for k in weekly["kpis"] if k["key"] == "target_completion_rate")
        conclusion = weekly["conclusions"][0]
        check(card["display"] != "—", "周报目标完成率卡片不应为 —")
        check(card["display"] in conclusion["text"],
              f"周报卡片 {card['display']} 与结论小数不一致：{conclusion['text']}")
        print("      rounding daily/weekly:", daily_display, card["display"])
    finally:
        client.close()


def t32_missing_source_failure():
    """缺源文件：批次失败、错误含文件名与影响指标、失败质量报告可展示、无新结论。"""
    with tempfile.TemporaryDirectory(prefix="mvp-missing-") as tmp:
        source = Path(tmp) / "source"
        shutil.copytree(SOURCE_DIR, source)
        (source / "leads.csv").unlink()
        db = Path(tmp) / "missing_source.db"
        result = run_batch(
            DEMO_DATE, trigger_type="selftest", triggered_by="selftest",
            db_path=db, source_dir=source,
        )
        check(result["status"] == "failed", f"缺源文件时批次应为 failed，实际 {result['status']}")
        error = result.get("error") or ""
        check("leads.csv" in error, f"批次错误未写出缺失文件：{error}")
        for metric in ("新增线索数", "有效线索数", "线索转化率", "线索异常风险"):
            check(metric in error, f"批次错误未写出影响指标 {metric}：{error}")
        failure = result.get("failure") or {}
        check(failure.get("kind") == "missing_source", "缺少结构化的 missing_source 失败信息")
        check(any(item.get("file") == "leads.csv" for item in failure.get("missing_files", [])),
              "失败信息缺少 leads.csv 明细")

        conn = sqlite3.connect(db)
        conn.row_factory = sqlite3.Row
        try:
            batch = conn.execute(
                "SELECT status, error FROM batch_runs WHERE batch_id=?", (result["batch_id"],)
            ).fetchone()
            check(batch is not None and batch["status"] == "failed", "batch_runs 状态不是 failed")
            check("leads.csv" in (batch["error"] or ""), "batch_runs.error 未保留缺失文件名")

            row = conn.execute(
                "SELECT report_json FROM quality_reports WHERE as_of=?", (DEMO_DATE.isoformat(),)
            ).fetchone()
            check(row is not None, "缺源文件时未落库可展示的失败质量报告")
            report = json.loads(row["report_json"])
            check(report["status"] == "fail", "失败质量报告状态应为 fail")
            check("leads.csv" in report.get("missing_sources", []), "质量报告 missing_sources 缺 leads.csv")
            labels = " ".join(m.get("label", "") for m in report.get("affected_metrics", []))
            check("新增线索数" in labels and "线索转化率" in labels,
                  f"质量报告影响指标不完整：{labels}")
            check(report.get("batch", {}).get("steps"), "失败质量报告缺少批次步骤")

            # 业务事务应整体回滚，不生成任何当日确定性结论。
            for table in ("metric_snapshots", "risks", "daily_briefs", "weekly_reports", "notifications"):
                count = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                check(count == 0, f"缺源文件失败后 {table} 不应有数据，实际 {count}")
        finally:
            conn.close()
        print("      missing source:", "leads.csv; affected 新增线索数/线索转化率; no conclusions")


def main() -> int:
    started = time.perf_counter()
    proc, base, port = start_server()
    print(f"临时服务：{base}　数据库：{DB_PATH.name}")
    passed = []
    failed = []
    clients = []
    try:
        gm = login_as(base, "u_gm")
        sm = login_as(base, "u_sm")
        sd = login_as(base, "u_sd")
        clients = [gm, sm, sd]
        scenarios = [
            ("01 健康检查 / 无移动端", lambda: t01_health_and_no_mobile(base)),
            ("02 未登录访问拦截", lambda: t02_unauth_guards(base)),
            ("03 四个演示账号与脱敏", lambda: t03_demo_accounts_masked(base)),
            ("04 登录 Cookie 安全属性", lambda: t04_login_cookie(base)),
            ("05 今日简报 6 指标 + 3 结论", lambda: t05_daily_brief(gm)),
            ("06 周报与上周对比", lambda: t06_weekly(gm)),
            ("07 12 项指标与下钻", lambda: t07_metrics_drill(gm)),
            ("08 数据质量与批次步骤", lambda: t08_quality(gm)),
            ("09 规则权限与参数", lambda: t09_rules(gm, sm)),
            ("10 风险清单对照 expected_risks", lambda: t10_risk_list_expected(gm)),
            ("11 四角色行级过滤", lambda: t11_row_level_scope(base)),
            ("12 风险详情证据链", lambda: t12_risk_detail(gm)),
            ("13 认领 → 关闭真实闭环", lambda: t13_claim_close(base, sm)),
            ("14 部门经理转派 403", lambda: t14_sm_transfer_forbidden(sm)),
            ("15 转派与原负责人关注权", lambda: t15_transfer_flow(gm)),
            ("16 手动升级", lambda: t16_escalate(gm)),
            ("17 规则启停/调参与权限", lambda: t17_rule_admin(gm, sd)),
            ("18 站内通知已读", lambda: t18_notifications(sm)),
            ("19 日报 Excel 导出", lambda: t19_export_brief(gm)),
            ("20 风险 Excel 行级过滤与脱敏", lambda: t20_export_risks_filtered(gm, sm)),
            ("21 打印版日报水印", lambda: t21_print(gm)),
            ("22 同日批处理重跑 3 次幂等", lambda: t22_batch_idempotent(gm)),
            ("23 批处理接口权限", lambda: t23_batch_permission(base, sm)),
            ("24 审计动作与越权记录", lambda: t24_audit(gm, sm)),
            ("25 OpenAPI 与 404", lambda: t25_openapi_and_404(base)),
            ("26 日报/周报计划发布时间", lambda: t26_report_publication_times(gm)),
            ("27 周报 Top3 按周区间过滤", lambda: t27_weekly_top3_scope(gm)),
            ("28 证据计数与可见证据一致", lambda: t28_evidence_consistency(gm)),
            ("29 列表/详情惰性超时升级幂等", lambda: t29_lazy_escalation(gm)),
            ("30 事业部二级升级到总经理", lambda: t30_bu_second_level_to_gm(base)),
            ("31 卡片与结论四舍五入一致", lambda: t31_card_conclusion_rounding(base)),
            ("32 缺少源文件错误与回滚", lambda: t32_missing_source_failure()),
        ]
        for name, func in scenarios:
            try:
                func()
            except Exception as exc:
                failed.append((name, str(exc)))
                print(f"[FAIL] {name}: {exc}")
            else:
                passed.append(name)
                print(f"[PASS] {name}")
    finally:
        for client in clients:
            client.close()
        stop_server(proc)
        for suffix in ("", "-wal", "-shm"):
            path = Path(str(DB_PATH) + suffix)
            if path.exists():
                path.unlink()
    elapsed = time.perf_counter() - started
    print("\n结果:", json.dumps({
        "passed": len(passed), "failed": len(failed), "total": len(passed) + len(failed),
        "elapsed_seconds": round(elapsed, 2), "port": port,
    }, ensure_ascii=False))
    if failed:
        print("失败项:")
        for name, message in failed:
            print(f"- {name}: {message}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
