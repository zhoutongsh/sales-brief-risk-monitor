"""Twenty independently designed end-to-end scenarios.

S01-S20 map to docs/测试用例清单.md. A failing Sxx is a real product deviation,
not a test harness error.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from recompute import DEMO_DATE, add_workdays, compute_metrics, compute_risks, half_up

TZ = timezone(timedelta(hours=8))
WATERMARK = "演示数据·禁止外传"
PHONE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")


def _daily(env) -> dict:
    response = env.get("/api/brief/daily")
    assert response.status_code == 200, response.text
    return response.json()


def _weekly(env, week: str = "this") -> dict:
    response = env.get(f"/api/brief/weekly?week={week}")
    assert response.status_code == 200, response.text
    return response.json()


def _risk_id(env, rule_id: str, object_id: str) -> str:
    row = env.one("SELECT id FROM risks WHERE rule_id=? AND object_id=?", (rule_id, object_id))
    assert row, (rule_id, object_id)
    return row["id"]


def _detail(env, risk_id: str) -> dict:
    response = env.get(f"/api/risks/{risk_id}")
    assert response.status_code == 200, response.text
    return response.json()


# ---------------------------------------------------------------------------
# S01-S04: four roles, daily and weekly scopes
# ---------------------------------------------------------------------------

def test_s01_gm_daily_brief(env):
    env.login("gm")
    brief = _daily(env)
    assert brief["date"] == "2026-10-24"
    assert brief["scope"]["name"] == "全公司"
    assert brief["scope"]["role"] == "gm"
    assert len(brief["kpis"]) == 6
    assert len(brief["conclusions"]) == 3
    assert len(brief["top_risks"]) == 3
    assert [r["id"] for r in brief["top_risks"]] == [
        "R-20261024-002", "R-20261024-003", "R-20261024-004",
    ]
    assert brief["risk_stats"]["total"] == 10
    assert brief["risk_stats"]["active"] == 10
    assert brief["risk_stats"]["red_active"] == 6
    assert "08:30" in brief["publish_schedule"]
    assert brief["data_quality"]["status"] in {"pass", "warn"}
    assert brief["published_at"]


def test_s02_sd_daily_and_weekly_scope(env):
    env.login("sd")
    brief = _daily(env)
    assert brief["scope"]["name"] == "全公司"
    assert brief["scope"]["role"] == "sd"
    assert brief["risk_stats"]["total"] == 10
    weekly = _weekly(env)
    assert weekly["scope"]["name"] == "全公司"
    assert weekly["week_start"] == "2026-10-19"
    assert weekly["data_through"] == "2026-10-24"
    assert weekly["partial"] is True
    assert len(weekly["kpis"]) == 6
    assert {row["id"] for row in weekly["ranking"]} == {"E100", "S200", "N300", "W400"}


def test_s03_bu_daily_and_weekly_scope(env):
    env.login("bu")
    brief = _daily(env)
    assert brief["scope"]["name"] == "华东事业部"
    assert brief["risk_stats"]["total"] == 6
    assert brief["risk_stats"]["red_active"] == 3
    assert brief["top_risks"][0]["id"] == "R-20261024-002"
    weekly = _weekly(env)
    assert weekly["scope"]["name"] == "华东事业部"
    assert {row["id"] for row in weekly["ranking"]} == {"E110", "E120", "E130", "E140"}
    items = env.get("/api/risks?page_size=100").json()["items"]
    assert len(items) == 6


def test_s04_sm_daily_and_weekly_scope(env):
    env.login("sm")
    brief = _daily(env)
    assert brief["scope"]["name"] == "华东一部"
    assert brief["risk_stats"]["total"] == 1
    assert len(brief["top_risks"]) == 1
    assert brief["top_risks"][0]["id"] == "R-20261024-008"
    weekly = _weekly(env)
    assert weekly["scope"]["name"] == "华东一部"
    assert weekly["partial"] is True
    assert {row["id"] for row in weekly["ranking"]} == {"u_sm", "u_e101", "u_e102", "u_e103"}
    items = env.get("/api/risks?page_size=100").json()["items"]
    assert [item["id"] for item in items] == ["R-20261024-008"]


# ---------------------------------------------------------------------------
# S05-S06: weekly this/last and risk-period correctness
# ---------------------------------------------------------------------------

def test_s05_gm_weekly_this_and_last(env):
    env.login("gm")
    this = _weekly(env, "this")
    last = _weekly(env, "last")
    assert this["week_start"] == "2026-10-19"
    assert this["week_end"] == "2026-10-25"
    assert this["data_through"] == "2026-10-24"
    assert this["edition"] == "周中版"
    assert last["week_start"] == "2026-10-12"
    assert last["week_end"] == "2026-10-18"
    assert last["data_through"] == "2026-10-18"
    assert last["edition"] == "上周正式版"
    assert this["comparison"]["current"]["start"] == "2026-10-19"
    assert this["comparison"]["previous"]["start"] == "2026-10-12"
    assert last["comparison"]["current"]["start"] == "2026-10-12"
    assert last["comparison"]["previous"]["start"] == "2026-10-05"
    assert this["risk_closure"]["new_count"] == 10
    assert last["risk_closure"]["new_count"] == 0
    assert len(this["daily"]) == 6
    assert len(last["daily"]) == 7
    assert len(this["kpis"]) == len(last["kpis"]) == 6


def test_s06_weekly_last_must_not_include_current_risks(env):
    """S06: last week's report must use risks from 10/12-10/18, not today's risks."""
    env.login("gm")
    last = _weekly(env, "last")
    ids = [risk["id"] for risk in last["top_risks"]]
    actual_periods = {
        row["id"]: row["as_of"] for row in env.rows("SELECT id, as_of FROM risks")
    }
    wrong = [risk_id for risk_id in ids if actual_periods.get(risk_id, "") > "2026-10-18"]
    assert not wrong, (
        "上周周报（10/12-10/18）的风险 Top3 混入本周风险：%s；"
        "同时 risk_closure.new_count=%s"
        % (wrong, last["risk_closure"]["new_count"])
    )


# ---------------------------------------------------------------------------
# S07-S12: risk handling
# ---------------------------------------------------------------------------

def test_s07_red_risk_four_hour_claim(env):
    risk_id = _risk_id(env, "target_gap", "E140")
    env.login("gm")
    data = _detail(env, risk_id)
    assert data["level"] == "red"
    created = datetime.fromisoformat(data["created_at"])
    claim_due = datetime.fromisoformat(data["sla"]["claim_due_at"])
    close_due = datetime.fromisoformat(data["sla"]["close_due_at"])
    assert claim_due - created == timedelta(hours=4)
    assert close_due == add_workdays(created, 3)
    response = env.post(f"/api/risks/{risk_id}/claim", json={"plan": "今天完成客户高层沟通"})
    assert response.status_code == 200
    claimed = response.json()
    assert claimed["status"] == "processing"
    assert claimed["sla"]["claim_state"] == "claimed"
    assert claimed["claimed_by"]["user_id"] == "u_gm"
    assert any(event["event_type"] == "claimed" for event in claimed["timeline"])
    assert env.post(f"/api/risks/{risk_id}/claim", json={}).status_code == 409
    audit = env.get("/api/audit?action=risk.claim").json()
    assert audit["total"] >= 1


def test_s08_timeout_escalation_by_batch(env):
    risk_id = _risk_id(env, "target_gap", "E130")
    env.execute(
        "UPDATE risks SET claim_due_at=?, escalation_level=0, escalated_to=NULL WHERE id=?",
        ("2026-10-24T00:00:00+08:00", risk_id),
    )
    result = env.run_batch()
    assert result["status"] == "success"
    assert result["stats"]["escalation"]["escalated"] >= 1
    env.login("gm")
    data = _detail(env, risk_id)
    assert data["escalation"]["level"] == 1
    assert data["escalation"]["escalated_to"]["user_id"] == "u_bu"
    assert any(
        event["event_type"] == "escalated" and event["payload"].get("automatic") is True
        for event in data["timeline"]
    )


def test_s09_lazy_escalation_on_risk_api(env):
    """S09: promised lazy check on opening risk center/detail (README, plan 4.1)."""
    risk_id = _risk_id(env, "target_gap", "E130")
    env.execute(
        "UPDATE risks SET claim_due_at=?, escalation_level=0, escalated_to=NULL WHERE id=?",
        ("2026-10-24T00:00:00+08:00", risk_id),
    )
    env.login("gm")
    env.get("/api/risks?page_size=100")
    detail = _detail(env, risk_id)
    assert detail["escalation"]["level"] == 1, (
        "认领时限已过，打开风险列表/详情后应只升 1 级（最多两级，列表加详情不能连升两次）"
    )
    auto_events = [
        event for event in detail["timeline"]
        if event["event_type"] == "escalated" and event["payload"].get("automatic") is True
    ]
    assert len(auto_events) == 1, "同一认领时限内重复打开了列表/详情，自动升级事件应只有 1 条"


def test_s10_transfer_changes_owner_and_sla(env):
    risk_id = _risk_id(env, "target_gap", "E140")
    env.login("bu")
    before = datetime.now(TZ)
    response = env.post(
        f"/api/risks/{risk_id}/transfer",
        json={"to_user_id": "u_ea3", "note": "华东三部更熟悉客户，转派跟进"},
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["owner"]["user_id"] == "u_ea3"
    assert data["status"] == "pending"
    assert data["escalation"]["level"] == 0
    assert any(w["user_id"] == "u_ea4" for w in data["watchers"])
    claim_due = datetime.fromisoformat(data["sla"]["claim_due_at"])
    assert timedelta(hours=4) - timedelta(seconds=120) <= claim_due - before <= timedelta(hours=4) + timedelta(seconds=120)
    assert any(event["event_type"] == "transferred" for event in data["timeline"])
    assert env.one(
        "SELECT COUNT(*) AS c FROM notifications WHERE user_id='u_ea3' AND title LIKE '%转派%'"
    )["c"] >= 1
    env.login("gm")
    audit = env.get("/api/audit?action=risk.transfer").json()
    assert audit["total"] >= 1


def test_s11_department_manager_transfer_forbidden(env):
    risk_id = _risk_id(env, "churn", "OPP-2026-0001")
    env.login("sm")
    response = env.post(
        f"/api/risks/{risk_id}/transfer",
        json={"to_user_id": "u_e101", "note": "转给代表跟进客户"},
    )
    assert response.status_code == 403
    detail = _detail(env, risk_id)
    assert detail["allowed_actions"]["transfer"] is False


def test_s12_close_and_state_machine_trail(env):
    risk_id = _risk_id(env, "target_gap", "E130")
    env.login("gm")
    assert env.post(f"/api/risks/{risk_id}/close", json={
        "result_status": "closed", "reason": "客户预算调整", "action": "重新排期", "result": "下月签约",
    }).status_code == 409
    assert env.post(f"/api/risks/{risk_id}/claim", json={"plan": "拉通重点商机"}).status_code == 200
    assert env.post(f"/api/risks/{risk_id}/close", json={
        "result_status": "closed", "reason": "太短", "action": "跟进", "result": "完成",
    }).status_code == 400
    response = env.post(f"/api/risks/{risk_id}/close", json={
        "result_status": "closed",
        "reason": "客户预算调整，目标顺延至下月",
        "action": "重新排期并逐单确认签约日期",
        "result": "已与客户确认 11 月首周签约",
    })
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["status"] == "closed"
    assert data["close"]["result_status"] == "closed"
    assert data["allowed_actions"]["claim"] is False
    assert data["allowed_actions"]["close"] is False
    event_types = [event["event_type"] for event in data["timeline"]]
    assert event_types[0] == "created"
    assert "claimed" in event_types and "closed" in event_types
    listing = env.get("/api/risks?page_size=100").json()["items"]
    assert listing[-1]["id"] == risk_id


# ---------------------------------------------------------------------------
# S13-S16: missing data, export, print, rules
# ---------------------------------------------------------------------------

def test_s13_missing_data_is_null_not_zero(env):
    env.login("gm")
    payload = env.get("/api/metrics?date=2026-10-23").json()
    assert payload["data_available"] is False
    assert len(payload["metrics"]) == 12
    assert all(item["value"] is None and item["display"] == "—" for item in payload["metrics"])
    assert payload["message"]
    assert env.get("/api/brief/daily?date=2026-10-23").status_code == 404
    assert env.get("/api/quality?date=2026-10-23").status_code == 404


def test_s14_real_xlsx_exports_with_watermark_and_scope(env):
    env.login("gm")
    brief_response = env.get("/api/export/brief.xlsx")
    assert brief_response.status_code == 200 and brief_response.content[:2] == b"PK"
    assert brief_response.headers["x-demo-watermark"] == "demo-data-do-not-distribute"
    brief_sheet = load_workbook(io.BytesIO(brief_response.content)).active
    brief_text = "\n".join(
        str(value) for row in brief_sheet.iter_rows(values_only=True) for value in row if value is not None
    )
    assert WATERMARK in str(brief_sheet["A1"].value)
    assert WATERMARK in (brief_sheet.oddHeader.center.text or "")
    assert "全公司" in brief_text

    risk_response = env.get("/api/export/risks.xlsx")
    risk_sheet = load_workbook(io.BytesIO(risk_response.content)).active
    risk_text = "\n".join(
        str(value) for row in risk_sheet.iter_rows(values_only=True) for value in row if value is not None
    )
    assert WATERMARK in risk_text
    assert not PHONE.search(risk_text)
    assert len(list(risk_sheet.iter_rows(min_row=5, values_only=True))) == 10

    env.login("sd")
    sd_sheet = load_workbook(io.BytesIO(env.get("/api/export/risks.xlsx").content)).active
    assert len(list(sd_sheet.iter_rows(min_row=5, values_only=True))) == 10

    env.login("bu")
    bu_sheet = load_workbook(io.BytesIO(env.get("/api/export/risks.xlsx").content)).active
    assert len(list(bu_sheet.iter_rows(min_row=5, values_only=True))) == 6
    env.login("sm")
    sm_sheet = load_workbook(io.BytesIO(env.get("/api/export/risks.xlsx").content)).active
    sm_rows = list(sm_sheet.iter_rows(min_row=5, values_only=True))
    assert len(sm_rows) == 1 and sm_rows[0][0] == "R-20261024-008"
    assert "138****" in str(sm_rows[0][6])
    env.login("gm")
    assert env.get("/api/audit?action=export.xlsx").json()["total"] >= 3


def test_s15_print_daily_brief(env):
    env.login("gm")
    response = env.get("/api/print/brief")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    html = response.text
    assert WATERMARK in html
    assert "window.print()" in html
    assert "三条结论" in html and "全公司" in html
    assert "2026-10-24" in html
    assert env.get("/api/audit?action=print.brief").json()["total"] >= 1


def test_s16_rules_edit_only_by_gm_and_audited(env):
    env.login("gm")
    rules = env.get("/api/rules").json()
    assert rules["total"] == 6 and rules["can_edit"] is True
    ok = env.patch("/api/rules/target_gap", json={"params": {"yellow_gap_pp": 2}})
    assert ok.status_code == 200
    assert ok.json()["params"] == {"yellow_gap_pp": 2, "orange_gap_pp": 8, "red_gap_pp": 15}
    assert env.patch("/api/rules/target_gap", json={"params": {"yellow_gap_pp": 9}}).status_code == 400
    assert env.patch("/api/rules/target_gap", json={"params": {"unknown": 1}}).status_code == 400
    disabled = env.patch("/api/rules/churn", json={"enabled": False})
    assert disabled.status_code == 200 and disabled.json()["enabled"] is False
    audit = env.get("/api/audit?action=rule.update").json()
    assert audit["total"] >= 2

    env.login("sd")
    assert env.get("/api/rules").json()["can_edit"] is False
    assert env.patch("/api/rules/target_gap", json={"params": {"yellow_gap_pp": 2}}).status_code == 403


# ---------------------------------------------------------------------------
# S17-S20: notifications, unauthorized access, rerun, publish timestamps
# ---------------------------------------------------------------------------

def test_s17_notifications_unread_and_read_all(env):
    env.login("gm")
    payload = env.get("/api/notifications").json()
    assert payload["total"] == 5
    assert payload["unread"] == 5
    assert any("简报" in item["title"] for item in payload["items"])
    first = payload["items"][0]["id"]
    read = env.post(f"/api/notifications/{first}/read")
    assert read.status_code == 200 and read.json()["is_read"] is True
    assert env.get("/api/notifications").json()["unread"] == 4
    assert env.get("/api/notifications?unread=true").json()["total"] == 4
    read_all = env.post("/api/notifications/read-all")
    assert read_all.status_code == 200 and read_all.json()["unread"] == 0
    assert env.get("/api/notifications?unread=true").json()["total"] == 0

    env.login("sm")
    assert env.post(f"/api/notifications/{first}/read").status_code == 404


def test_s18_unauthorized_access_is_blocked_and_audited(env):
    assert env.get("/api/risks").status_code == 401
    assert env.get("/api/export/brief.xlsx").status_code == 401
    env.login("sm")
    outsider = env.one("SELECT id FROM risks WHERE object_id='E140'")["id"]
    assert env.get(f"/api/risks/{outsider}").status_code == 403
    assert env.get("/api/audit").status_code == 403
    assert env.post("/api/admin/batch/run", json={}).status_code == 403
    env.login("gm")
    denied = env.get("/api/audit?action=access.denied&limit=100").json()
    assert denied["total"] >= 2


def test_s19_same_day_rerun_is_idempotent(env):
    env.login("gm")
    counts_before = {
        table: env.one(f"SELECT COUNT(*) AS c FROM {table}")["c"]
        for table in ["risks", "risk_evidence", "risk_events", "daily_briefs",
                      "weekly_reports", "notifications", "metric_snapshots"]
    }
    risk_before = env.rows(
        "SELECT id, level, status, owner_user_id, amount, claim_due_at, close_due_at "
        "FROM risks ORDER BY id"
    )
    metrics_before = env.rows(
        "SELECT scope_id, metric_key, value FROM metric_snapshots ORDER BY scope_id, metric_key"
    )
    brief_before = env.get("/api/brief/daily").json()
    for _ in range(2):
        result = env.run_batch()
        assert result["status"] == "success", result.get("error")
        assert result["stats"]["risks"]["skipped"] == 10
        assert result["stats"]["risks"]["created"] == 0
        assert result["stats"]["risks"]["merged"] == 0
    counts_after = {
        table: env.one(f"SELECT COUNT(*) AS c FROM {table}")["c"]
        for table in counts_before
    }
    assert counts_after == counts_before
    assert env.rows(
        "SELECT id, level, status, owner_user_id, amount, claim_due_at, close_due_at "
        "FROM risks ORDER BY id"
    ) == risk_before
    assert env.rows(
        "SELECT scope_id, metric_key, value FROM metric_snapshots ORDER BY scope_id, metric_key"
    ) == metrics_before
    brief_after = env.get("/api/brief/daily").json()
    assert [k["value"] for k in brief_after["kpis"]] == [k["value"] for k in brief_before["kpis"]]
    assert [r["id"] for r in brief_after["top_risks"]] == [r["id"] for r in brief_before["top_risks"]]
    assert brief_after["risk_stats"] == brief_before["risk_stats"]
    assert env.one("SELECT COUNT(*) AS c FROM batch_runs")["c"] == 3


def test_s20_publish_timestamps(env):
    """S20: daily 08:30, weekly this Saturday 09:00, weekly last Monday 09:00."""
    env.login("gm")
    daily = _daily(env)
    this = _weekly(env, "this")
    last = _weekly(env, "last")
    errors = []
    expected_daily = "2026-10-24T08:30:00+08:00"
    if daily["published_at"] != expected_daily:
        errors.append(f"日报 published_at={daily['published_at']}，应为 {expected_daily}")
    expected_this = "2026-10-24T09:00:00+08:00"
    if this["published_at"] != expected_this:
        errors.append(f"本周周报 published_at={this['published_at']}，应为 {expected_this}（周六 09:00 周中版）")
    expected_last = "2026-10-19T09:00:00+08:00"
    if last["published_at"] != expected_last:
        errors.append(f"上周周报 published_at={last['published_at']}，应为 {expected_last}（周一 09:00 正式版）")
    assert not errors, "；".join(errors)
