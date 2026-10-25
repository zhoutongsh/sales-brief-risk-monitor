"""Real xlsx exports and printable daily brief."""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook

WATERMARK = "演示数据·禁止外传"
PHONE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")


def _xlsx(response):
    assert response.status_code == 200, response.text
    assert response.content[:2] == b"PK"
    assert response.headers["x-demo-watermark"] == "demo-data-do-not-distribute"
    return load_workbook(io.BytesIO(response.content))


def _sheet_text(ws) -> str:
    parts = []
    for row in ws.iter_rows(values_only=True):
        parts.extend(str(value) for value in row if value is not None)
    return "\n".join(parts)


def test_brief_xlsx_is_real_and_watermarked(env):
    env.login("gm")
    response = env.get("/api/export/brief.xlsx")
    workbook = _xlsx(response)
    sheet = workbook.active
    assert WATERMARK in str(sheet["A1"].value)
    assert WATERMARK in str(sheet["A3"].value)
    assert WATERMARK in (sheet.oddHeader.center.text or "")
    assert WATERMARK in (sheet.oddFooter.center.text or "")
    text = _sheet_text(sheet)
    assert "全公司" in text
    assert "三条结论" in text
    assert "风险 Top3" in text
    assert response.headers["content-disposition"].startswith("attachment;")
    assert "brief_20261024_u_gm.xlsx" in response.headers["content-disposition"]

    audit = env.get("/api/audit?action=export.xlsx").json()
    assert audit["total"] >= 1


def test_brief_xlsx_scope_differs_by_role(env):
    env.login("bu")
    workbook = _xlsx(env.get("/api/export/brief.xlsx"))
    text = _sheet_text(workbook.active)
    assert "华东事业部" in text
    assert "全公司" not in text.split("数据范围：")[1].split("　")[0]


def test_risk_xlsx_is_scoped_masked_and_watermarked(env):
    env.login("gm")
    gm_book = _xlsx(env.get("/api/export/risks.xlsx"))
    gm_sheet = gm_book.active
    gm_text = _sheet_text(gm_sheet)
    assert WATERMARK in str(gm_sheet["A1"].value)
    assert WATERMARK in (gm_sheet.oddFooter.center.text or "")
    assert not PHONE.search(gm_text), "raw phone leaked into xlsx"
    gm_rows = list(gm_sheet.iter_rows(min_row=5, values_only=True))
    assert len(gm_rows) == 10
    assert all(row[0].startswith("R-20261024-") for row in gm_rows)

    env.login("bu")
    bu_sheet = _xlsx(env.get("/api/export/risks.xlsx")).active
    bu_rows = list(bu_sheet.iter_rows(min_row=5, values_only=True))
    assert len(bu_rows) == 6
    assert all("R-20261024" in row[0] for row in bu_rows)

    env.login("sm")
    sm_sheet = _xlsx(env.get("/api/export/risks.xlsx")).active
    sm_rows = list(sm_sheet.iter_rows(min_row=5, values_only=True))
    assert len(sm_rows) == 1
    assert sm_rows[0][0] == "R-20261024-008"
    assert "138****" in str(sm_rows[0][6])


def test_risk_export_level_filter(env):
    env.login("gm")
    sheet = _xlsx(env.get("/api/export/risks.xlsx?level=red")).active
    rows = list(sheet.iter_rows(min_row=5, values_only=True))
    assert len(rows) == 6
    assert all(row[1] == "红色" for row in rows)


def test_print_brief_is_watermarked_html(env):
    env.login("gm")
    response = env.get("/api/print/brief")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    html = response.text
    assert WATERMARK in html
    assert "window.print()" in html
    assert "全公司" in html
    assert "三条结论" in html
    audit = env.get("/api/audit?action=print.brief").json()
    assert audit["total"] >= 1
